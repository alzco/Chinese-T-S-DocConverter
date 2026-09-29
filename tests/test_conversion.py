import io
import unittest
from zipfile import ZipFile

from docx import Document
from lxml import etree
from opencc import OpenCC

from document_converter import DocumentConverter, W
from opencc_converter import CustomOpenCC, DATA_DIR, UPSTREAM_CONFIGS, STANDARD_CONFIGS
from resources import SOURCE, dictionary_bundle, DOWNLOADS
import hashlib
import json


class ConversionTests(unittest.TestCase):
    def test_upstream_files_unchanged(self):
        for name, digest in SOURCE['files'].items():
            with self.subTest(file=name):
                self.assertEqual(hashlib.sha256((DATA_DIR / name).read_bytes()).hexdigest(), digest)

    def test_default_is_upstream_standard(self):
        text = '为众人说《西厢记》：两晋异体字、绿色记录、启用线索。'
        self.assertEqual(CustomOpenCC().convert(text), '爲衆人説《西厢記》：兩晋异體字、緑色記録、啓用綫索。')
        self.assertNotEqual(CustomOpenCC().convert(text), OpenCC('s2t').convert(text))
        # The removed document's terminology overrides must no longer run.
        self.assertEqual(CustomOpenCC().convert('資訊、專案、劇碼'), '資訊、專案、劇碼')

    def test_all_upstream_configs_and_normalization(self):
        samples = ['為眾人說話，綠色記錄。', '皇后以后，干净干部，白云云云。', '\ufa19']
        for config, filename in UPSTREAM_CONFIGS.items():
            with self.subTest(config=config):
                reference = OpenCC(str(DATA_DIR / filename))
                for text in samples:
                    self.assertEqual(CustomOpenCC(config).convert(text), reference.convert(text))
        self.assertEqual(CustomOpenCC().convert('\ufa19'), '神')
        self.assertEqual(CustomOpenCC('t2gov').convert('綠色記錄'), '緑色記録')

    def test_standard_opencc_directions(self):
        for config in STANDARD_CONFIGS:
            with self.subTest(config=config):
                text = '為眾人說話，信息、项目、绿色。'
                self.assertEqual(CustomOpenCC(config).convert(text), OpenCC(config).convert(text))

    def test_download_bundle_contains_all_dependencies(self):
        with ZipFile(io.BytesIO(dictionary_bundle())) as archive:
            for filename in DOWNLOADS.values():
                self.assertEqual(archive.read('tonggui/' + filename), (DATA_DIR / filename).read_bytes())
            self.assertIn('tonggui/LICENSE', archive.namelist())
            self.assertIn('tonggui/SOURCE.json', archive.namelist())
            def verify_dictionary(dictionary):
                if 'file' in dictionary:
                    self.assertIn('tonggui/' + dictionary['file'], archive.namelist())
                for child in dictionary.get('dicts', []):
                    verify_dictionary(child)
            for filename in UPSTREAM_CONFIGS.values():
                config = json.loads(archive.read('tonggui/' + filename))
                for stage in config.get('normalization', []) + config['conversion_chain']:
                    verify_dictionary(stage['dict'])

    def test_dictionary_validation(self):
        from opencc_converter import validate_dictionary
        for invalid in [[], {'': '空'}, {'字': 3}]:
            with self.assertRaises(ValueError):
                validate_dictionary(invalid)

    def test_custom_dictionary_and_plain_files(self):
        converter = DocumentConverter(custom_dict={'计算机': '信息处理设备'})
        for kind in ['txt', 'md']:
            self.assertEqual(converter.convert_file('计算机为众人服务'.encode(), kind), '信息處理設備爲衆人服務')
        self.assertEqual(converter.convert_txt_file('绿色'.encode('gbk')), '緑色')

    def test_docx_all_parts_formatting_and_cross_run_context(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run('皇').bold = True
        paragraph.add_run('后以后为众人说话').italic = True
        document.add_table(rows=1, cols=1).cell(0, 0).text = '绿色记录'
        document.sections[0].header.paragraphs[0].text = '启用线索'
        document.sections[0].footer.paragraphs[0].text = '西厢记'
        original = io.BytesIO()
        document.save(original)
        package = io.BytesIO()
        with ZipFile(original) as source, ZipFile(package, 'w') as target:
            for item in source.infolist():
                target.writestr(item, source.read(item.filename))
            for kind in ['footnote', 'endnote']:
                target.writestr(f'word/{kind}s.xml', f'<w:{kind}s xmlns:w="{W[1:-1]}"><w:{kind} w:id="1"><w:p><w:r><w:rPr><w:b/></w:rPr><w:t>为众人说话</w:t></w:r></w:p></w:{kind}></w:{kind}s>')
            target.writestr('word/media/untouched.bin', b'keep this image')
        output = DocumentConverter().convert_docx_file(package.getvalue())
        reread = Document(io.BytesIO(output))
        self.assertEqual(reread.paragraphs[0].text, '皇后以後爲衆人説話')
        self.assertTrue(reread.paragraphs[0].runs[0].bold)
        self.assertTrue(reread.paragraphs[0].runs[1].italic)
        self.assertEqual(reread.tables[0].cell(0, 0).text, '緑色記録')
        self.assertEqual(reread.sections[0].header.paragraphs[0].text, '啓用綫索')
        self.assertEqual(reread.sections[0].footer.paragraphs[0].text, '西厢記')
        with ZipFile(io.BytesIO(output)) as result:
            for kind in ['footnote', 'endnote']:
                root = etree.fromstring(result.read(f'word/{kind}s.xml'))
                self.assertEqual(''.join(root.itertext()), '爲衆人説話')
                self.assertIsNotNone(root.find('.//' + W + 'b'))
            self.assertEqual(result.read('word/media/untouched.bin'), b'keep this image')

    def test_cross_run_custom_length_changes_and_hyperlinks(self):
        root = etree.fromstring(f'<w:p xmlns:w="{W[1:-1]}"><w:r><w:t>计</w:t></w:r><w:hyperlink><w:r><w:t>算机</w:t></w:r></w:hyperlink><w:r><w:t>为众人</w:t></w:r></w:p>')
        converter = DocumentConverter(custom_dict={'计算机': '信息处理设备'})
        converter._convert_xml_paragraph(root)
        self.assertEqual(''.join(root.itertext()), '信息處理設備爲衆人')
        self.assertIsNotNone(root.find(W + 'hyperlink'))


if __name__ == '__main__':
    unittest.main()
